"""Deterministic SVG and editable PowerPoint rendering."""

from __future__ import annotations

import html
import json
from pathlib import Path

from .assets import resolve_node_asset
from .layout import LayoutScene, Point
from .models import SemanticScene

PALETTE = {
    "actor": ("#eff6ff", "#93c5fd", "#2563eb"), "service": ("#ffffff", "#cbd5e1", "#334155"),
    "database": ("#f0fdf4", "#86efac", "#16a34a"), "cache": ("#fff7ed", "#fdba74", "#ea580c"),
    "queue": ("#faf5ff", "#d8b4fe", "#9333ea"), "external": ("#f8fafc", "#cbd5e1", "#64748b"),
    "image": ("#fdf2f8", "#f9a8d4", "#db2777"), "asset": ("#fdf2f8", "#f9a8d4", "#db2777"),
    "note": ("#fffbeb", "#fcd34d", "#d97706"), "group": ("#f8fafc", "#94a3b8", "#475569"),
}


def _esc(value: str) -> str:
    return html.escape(value, quote=True)


def _route(points: tuple[Point, ...]) -> str:
    return "M " + " L ".join(f"{point.x:g} {point.y:g}" for point in points)


def render_svg(layout: LayoutScene, scene: SemanticScene) -> str:
    title = _esc(scene.intent.title if scene.intent and scene.intent.title else scene.id)
    purpose = _esc(scene.intent.purpose) if scene.intent and scene.intent.purpose else ""
    semantic = {node.id: node for node in scene.nodes}
    edges = []
    for edge in layout.edges:
        dash = ' stroke-dasharray="7 6"' if edge.kind in {"async", "event"} else ""
        label = ""
        if edge.label:
            point = edge.points[len(edge.points) // 2]
            label = f'<text x="{point.x + 8:g}" y="{point.y - 8:g}" font-family="Inter,system-ui" font-size="11" fill="#475569">{_esc(edge.label)}</text>'
        edges.append(f'<g data-edge-id="{_esc(edge.id)}"><path d="{_route(edge.points)}" fill="none" stroke="#64748b" stroke-width="2"{dash} marker-end="url(#arrow)"/>{label}</g>')
    nodes = []
    for node in layout.nodes:
        original = semantic[node.id]
        fill, stroke, accent = PALETTE.get(node.kind, PALETTE["service"])
        strong = original.importance == "primary" or (original.style and original.style.emphasis == "strong")
        shape_kind = original.metadata.get("shape")
        if shape_kind in {"ellipse", "diamond", "rectangle"}:
            cx, cy = node.x + node.width / 2, node.y + node.height / 2
            attrs = 'fill="#ffffff" stroke="#64748b" stroke-width="2"'
            if shape_kind == "ellipse":
                shape = f'<ellipse cx="{cx:g}" cy="{cy:g}" rx="{node.width/2:g}" ry="{node.height/2:g}" {attrs}/>'
            elif shape_kind == "diamond":
                shape = f'<polygon points="{cx:g},{node.y:g} {node.x+node.width:g},{cy:g} {cx:g},{node.y+node.height:g} {node.x:g},{cy:g}" {attrs}/>'
            else:
                shape = f'<rect x="{node.x:g}" y="{node.y:g}" width="{node.width:g}" height="{node.height:g}" rx="4" {attrs}/>'
            nodes.append(f'<g data-scene-id="{_esc(node.id)}" data-kind="{_esc(node.kind)}">{shape}<text x="{cx:g}" y="{cy+5:g}" text-anchor="middle" font-family="Inter,system-ui" font-size="14" fill="#0f172a">{_esc(node.label)}</text></g>')
            continue
        icon = resolve_node_asset(original)["glyph"]
        icon_x, icon_y = node.x + 18, node.y + node.height / 2 - 15
        nodes.append(
            f'<g data-scene-id="{_esc(node.id)}" data-kind="{_esc(node.kind)}">'
            f'<rect x="{node.x:g}" y="{node.y:g}" width="{node.width:g}" height="{node.height:g}" rx="16" fill="{fill}" stroke="{accent if strong else stroke}" stroke-width="{2.5 if strong else 1.5}"/>'
            f'<rect x="{icon_x:g}" y="{icon_y:g}" width="34" height="30" rx="9" fill="{accent}"/>'
            f'<text x="{icon_x + 17:g}" y="{icon_y + 19:g}" text-anchor="middle" font-family="Inter,system-ui" font-size="10" font-weight="700" fill="#fff">{_esc(icon)}</text>'
            f'<text x="{node.x + 62:g}" y="{node.y + node.height / 2 + 5:g}" font-family="Inter,system-ui" font-size="15" font-weight="{700 if strong else 600}" fill="#0f172a">{_esc(node.label)}</text></g>'
        )
    subtitle = f'<text x="64" y="58" font-family="Inter,system-ui" font-size="13" fill="#64748b">{purpose}</text>' if purpose else ""
    return f'''<svg xmlns="http://www.w3.org/2000/svg" width="{layout.width:g}" height="{layout.height:g}" viewBox="0 0 {layout.width:g} {layout.height:g}">
<defs><marker id="arrow" viewBox="0 0 10 10" refX="9" refY="5" markerWidth="7" markerHeight="7" orient="auto-start-reverse"><path d="M 0 0 L 10 5 L 0 10 z" fill="#64748b"/></marker></defs>
<rect width="100%" height="100%" fill="#f8fafc"/><text x="64" y="34" font-family="Inter,system-ui" font-size="22" font-weight="700" fill="#0f172a">{title}</text>{subtitle}{''.join(edges)}{''.join(nodes)}</svg>'''


def render_html(scene: SemanticScene, svg: str) -> str:
    from .inspector import inspector_html
    return inspector_html(scene, svg)


def write_pptx(layout: LayoutScene, scene: SemanticScene, output: str | Path) -> None:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
    from pptx.oxml.xmlchemy import OxmlElement
    from pptx.util import Inches, Pt

    presentation = Presentation()
    presentation.slide_width, presentation.slide_height = Inches(13.333), Inches(7.5)
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    background = slide.background.fill
    background.solid(); background.fore_color.rgb = _rgb("F8FAFC")
    title = slide.shapes.add_textbox(Inches(.55), Inches(.18), Inches(12.2), Inches(.4))
    title.text_frame.paragraphs[0].text = scene.intent.title if scene.intent and scene.intent.title else scene.id
    title.text_frame.paragraphs[0].font.size = Pt(20); title.text_frame.paragraphs[0].font.bold = True
    scale = min(12.333 / max(layout.width, 1), 6.3 / max(layout.height, 1))
    xoff, yoff = (13.333 - layout.width * scale) / 2, .85
    for edge in layout.edges:
        for index, (a, b) in enumerate(zip(edge.points, edge.points[1:])):
            connector = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(xoff + a.x * scale), Inches(yoff + a.y * scale), Inches(xoff + b.x * scale), Inches(yoff + b.y * scale))
            connector.line.color.rgb = _rgb("64748B"); connector.line.width = Pt(1.3)
            if edge.kind in {"async", "event"}:
                connector.line.dash_style = 4
            if index == len(edge.points) - 2:
                line = connector._element.spPr.get_or_add_ln()
                arrow = OxmlElement("a:tailEnd")
                arrow.set("type", "triangle")
                line.append(arrow)
        if edge.label:
            point = edge.points[len(edge.points) // 2]
            box = slide.shapes.add_textbox(Inches(xoff + (point.x + 8) * scale), Inches(yoff + (point.y - 24) * scale), Inches(.6), Inches(.25))
            box.text_frame.paragraphs[0].text = edge.label
            box.text_frame.paragraphs[0].font.size = Pt(10)
    by_id = {node.id: node for node in scene.nodes}
    for node in layout.nodes:
        original = by_id[node.id]
        fill, stroke, accent = [value.lstrip("#") for value in PALETTE.get(node.kind, PALETTE["service"])]
        x, y, w, h = xoff + node.x * scale, yoff + node.y * scale, node.width * scale, node.height * scale
        shape_type = {"ellipse": MSO_SHAPE.OVAL, "diamond": MSO_SHAPE.DIAMOND, "rectangle": MSO_SHAPE.RECTANGLE}.get(original.metadata.get("shape"), MSO_SHAPE.ROUNDED_RECTANGLE)
        shape = slide.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h))
        shape.fill.solid(); shape.fill.fore_color.rgb = _rgb(fill); shape.line.color.rgb = _rgb(stroke)
        shape.text_frame.clear(); shape.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
        paragraph = shape.text_frame.paragraphs[0]; paragraph.alignment = PP_ALIGN.CENTER; paragraph.text = node.label
        paragraph.font.size = Pt(max(8, min(14, h * 17))); paragraph.font.bold = original.importance == "primary"
    output = Path(output); output.parent.mkdir(parents=True, exist_ok=True); presentation.save(output)


def _rgb(value: str):
    from pptx.dml.color import RGBColor
    return RGBColor.from_string(value)
