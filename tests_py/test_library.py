import json
import xml.etree.ElementTree as ET
from fastapi.testclient import TestClient
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from scenelyr.api import app, ui_scenes
from scenelyr.models import SemanticScene
from scenelyr.storage import save_scene
from scenelyr.mcp_server import export_scene


def test_saved_layers_reopen_and_shape_exports(tmp_path, monkeypatch):
    monkeypatch.setenv('SCENELYR_DATA_DIR', str(tmp_path))
    scene = SemanticScene.model_validate({'id':'layer-test','nodes':[
        {'id':'a','label':'Decision','kind':'asset','metadata':{'shape':'diamond','sourceBounds':[1,2,30,40]}},
        {'id':'b','label':'End','kind':'asset','metadata':{'shape':'ellipse'}}],
        'edges':[{'id':'ab','from':'a','to':'b','label':'Yes'}]})
    paths = save_scene(scene)
    manifest=json.loads((paths['folder']/'assets/manifest.json').read_text())
    assert len(manifest['layers'])==2
    assert manifest['layers'][0]['sourceBounds']==[1,2,30,40]
    assert manifest['layers'][0]['position']['width']>0
    ET.parse(manifest['layers'][0]['file'])
    ui_scenes.clear()
    response=TestClient(app).get('/ui/scenes/layer-test')
    assert response.status_code==200
    assert response.json()['nodes'][0]['label']=='Decision'
    assert 'Diagram layers' in response.json()['inspector']
    assert '<polygon' in response.json()['svg'] and '<ellipse' in response.json()['svg']
    shapes=Presentation(paths['pptx']).slides[0].shapes
    assert any(s.text=='Yes' for s in shapes if s.has_text_frame)
    assert any(s.auto_shape_type==MSO_SHAPE.DIAMOND for s in shapes if s.shape_type==1)
    exported=json.loads(export_scene('layer-test',str(tmp_path/'exports')))
    assert json.loads(open(exported['json']).read())['id']=='layer-test'
